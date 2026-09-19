"""Generate the public, fictional course timetable sample with no participant data."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

workbook = Workbook()
sheet = workbook.active
sheet.title = '주차별 시간표'
sheet.append(['2026 과정 일정 샘플'])
sheet.merge_cells('A1:L1')
sheet.append(['4행의 열 이름을 유지하세요. 날짜는 YYYY-MM-DD 또는 M/D, 시간대는 HH:mm~HH:mm입니다. OT는 0주차, 휴강·합계는 제외됩니다.'])
sheet.merge_cells('A2:L2')
sheet.append([])
sheet.append(['주차', '주제(테마)', '구분', '날짜', '요일', '시간(h)', '시간대', '세션명', '세부 내용', '강사(안)', '산출물', '달력 표시용(자동)'])
rows = [
    ['OT', '오리엔테이션', 'OT', '2026-09-12', '토', 2, '10:00~12:00', '과정 안내 및 준비', '운영 방식과 학습 목표 소개', '담당 강사', '학습 계획서'],
    ['1주차', '기초 실습', '온라인1', '9/15', '화', 3, '19:00~22:00', 'AI 도구 이해', '도구 사용법과 기초 실습', '담당 강사', '실습 기록'],
    ['1주차', '기초 실습', '온라인2', '9/17', '목', 3, '19:00~22:00', '프로젝트 시작하기', '개발 환경 준비', '담당 강사', '프로젝트 초안'],
    ['1주차', '기초 실습', '오프라인', '9/19', '토', 5, '10:00~15:00', '팀 프로젝트 계획', '팀별 목표와 역할 정하기', '담당 강사', '팀 계획서'],
    ['휴강', '휴강 예시', '휴강', '9/24', '목', 0, '—', '휴강', '수업으로 등록되지 않습니다'],
    ['2주차', '프로젝트 실습', '온라인1', '9/29', '화', 3, '19:00~22:00', '문제 정의와 요구사항', '프로젝트 요구사항 정리', '담당 강사', '요구사항 문서'],
    ['합계', '', '정규 과정', '', '', 16, '', 'OT 포함 5개 수업'],
]
for row in rows:
    sheet.append(row)
for i in range(5, 5 + len(rows)):
    sheet.cell(i, 12, f'=C{i}&" · "&F{i}&"h · "&H{i}')
widths = [12, 24, 16, 18, 9, 12, 20, 32, 45, 18, 24, 40]
for col, width in enumerate(widths, 1):
    sheet.column_dimensions[get_column_letter(col)].width = width
for row in sheet:
    for cell in row:
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.font = Font(name='맑은 고딕', size=11, color='203C32')
    sheet.row_dimensions[row[0].row].height = 42
for cell in sheet[4]:
    cell.fill = PatternFill('solid', fgColor='244B3D')
    cell.font = Font(name='맑은 고딕', bold=True, color='FFFFFF')
sheet.freeze_panes = 'D5'
sheet.auto_filter.ref = f'A4:L{sheet.max_row}'
guide = workbook.create_sheet('작성 안내')
for row in [
    ['과정 일정 가져오기 안내'],
    ['1', '주차별 시간표 시트를 편집하고 파일을 선택하세요. 다른 시트는 가져오지 않습니다.'],
    ['2', '주차·날짜·시간대·세션명은 필수입니다. 주차는 OT 또는 1~52주차입니다. 수업은 최대 104개입니다.'],
    ['3', '연도가 없는 날짜는 제목의 연도 또는 화면의 기준 연도를 사용합니다. 연도를 넘기는 일정은 YYYY-MM-DD로 입력하세요.'],
    ['4', '요일을 입력했다면 날짜와 일치해야 합니다. 같은 날 시간이 겹치는 수업은 수정하세요.'],
    ['5', '휴강과 합계 행은 제외됩니다. 별도 멘토링 시트는 날짜·시간을 확정한 뒤 시간표 시트에 추가하세요.'],
    ['6', '미리보기에서 오류를 확인하고 편집에 반영하세요. 기존 일정과 운영 기간·주차가 교체됩니다.'],
    ['7', '마지막으로 일정 저장을 눌러야 서버에 저장됩니다. 파일을 선택하는 것만으로는 저장되지 않습니다.'],
    ['8', '강사·세부 내용·산출물은 일정 메모로 보관됩니다. 실제 강사 계정이나 출석 회차를 생성하지 않습니다.'],
]:
    guide.append(row)
guide.column_dimensions['A'].width = 15
guide.column_dimensions['B'].width = 110
for row in guide:
    guide.row_dimensions[row[0].row].height = 42
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical='center')
destination = Path(__file__).resolve().parents[1] / 'public/templates/course-schedule-template.xlsx'
destination.parent.mkdir(parents=True, exist_ok=True)
workbook.save(destination)
print(destination.name)
